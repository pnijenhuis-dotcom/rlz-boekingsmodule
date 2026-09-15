"""Cel-raster als gedeelde vorm voor .xls (xlrd), .xlsx (openpyxl) en JSON-fixtures (gouden set — een echte .xls is niet
herschrijfbaar zonder xlwt, dus de geanonimiseerde fixture is het raster zelf). Alleen niet-lege cellen; datums als
`datetime.date`, getallen als `Decimal` (financieel: nooit floats doorgeven aan de parsers)."""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import PurePosixPath
from typing import Any

Cel = str | Decimal | date | None


def _naar_cel(waarde: Any, *, xls_datum: bool = False, datemode: int = 0) -> Cel:
    if waarde is None:
        return None
    if isinstance(waarde, datetime):
        return waarde.date()
    if isinstance(waarde, date):
        return waarde
    if isinstance(waarde, bool):
        return Decimal(int(waarde))
    if isinstance(waarde, int):
        return Decimal(waarde)
    if isinstance(waarde, float):
        if xls_datum:
            import xlrd

            return xlrd.xldate.xldate_as_datetime(waarde, datemode).date()
        # POS-exports schrijven floats (86.81000000000001): op de cent afronden is hier de enige juiste lezing —
        # de bron rekent in centen.
        return Decimal(repr(waarde)).quantize(Decimal("0.01")) if abs(waarde) < 10**12 else Decimal(repr(waarde))
    if isinstance(waarde, Decimal):
        return waarde
    tekst = str(waarde)
    return tekst if tekst.strip() != "" else None


@dataclass
class Grid:
    """Rijen van {kolomindex: cel}; `naam` = bladnaam (meerdere bladen → de aanroeper kiest via `bladen`)."""

    naam: str
    rijen: list[dict[int, Cel]]
    bladen: dict[str, Grid] = field(default_factory=dict)

    def tekst(self, r: int, c: int) -> str | None:
        w = self.rijen[r].get(c) if 0 <= r < len(self.rijen) else None
        return w if isinstance(w, str) else None

    def getal(self, r: int, c: int) -> Decimal | None:
        w = self.rijen[r].get(c) if 0 <= r < len(self.rijen) else None
        if isinstance(w, Decimal):
            return w
        if isinstance(w, str):
            try:
                return Decimal(w.replace(",", ".").replace("€", "").strip())
            except InvalidOperation:
                return None
        return None

    def zoek(self, label: str, *, exact: bool = False) -> tuple[int, int] | None:
        """Eerste cel waarvan de tekst (gestript, hoofdletterongevoelig) met `label` begint (of gelijk is)."""
        doel = label.strip().lower()
        for r, rij in enumerate(self.rijen):
            for c in sorted(rij):
                w = rij[c]
                if isinstance(w, str):
                    s = w.strip().lower()
                    if (s == doel) if exact else s.startswith(doel):
                        return r, c
        return None

    def getallen_in_rij(self, r: int, *, kol_tot: int | None = None, kol_vanaf: int = 0) -> list[Decimal]:
        rij = self.rijen[r] if 0 <= r < len(self.rijen) else {}
        return [
            rij[c]
            for c in sorted(rij)
            if isinstance(rij[c], Decimal) and c >= kol_vanaf and (kol_tot is None or c < kol_tot)
        ]

    def als_json(self) -> list[list[list[Any]]]:
        """Serialiseerbare vorm voor fixtures: per rij [[kolom, waarde], …] (Decimal → str, date → ISO)."""
        uit: list[list[list[Any]]] = []
        for rij in self.rijen:
            cellen = []
            for c in sorted(rij):
                w = rij[c]
                cellen.append([c, w.isoformat() if isinstance(w, date) else str(w) if isinstance(w, Decimal) else w])
            uit.append(cellen)
        return uit


def grid_uit_json(naam: str, data: list[list[list[Any]]]) -> Grid:
    rijen: list[dict[int, Cel]] = []
    for cellen in data:
        rij: dict[int, Cel] = {}
        for c, w in cellen:
            if isinstance(w, str):
                try:
                    rij[int(c)] = (
                        date.fromisoformat(w) if len(w) == 10 and w[4] == "-" and w[7] == "-" else _decimal_of_tekst(w)
                    )
                except ValueError:
                    rij[int(c)] = _decimal_of_tekst(w)
            else:
                rij[int(c)] = _naar_cel(w)
        rijen.append(rij)
    return Grid(naam=naam, rijen=rijen)


def _decimal_of_tekst(w: str) -> Cel:
    # Fixture-waarden die uit een Decimal kwamen staan als "250.00"; tekst blijft tekst.
    if w and all(ch in "0123456789.-" for ch in w) and any(ch.isdigit() for ch in w):
        try:
            return Decimal(w)
        except InvalidOperation:
            return w
    return w


def lees_grid(bestandsnaam: str, inhoud: bytes) -> Grid:
    """Leest álle bladen; het teruggegeven Grid is het eerste blad mét `bladen` voor de rest."""
    suffix = PurePosixPath(bestandsnaam).suffix.lower()
    if suffix == ".xls":
        return _lees_xls(inhoud)
    if suffix == ".xlsx":
        return _lees_xlsx(inhoud)
    raise ValueError(f"Geen spreadsheet: {bestandsnaam}")


def _lees_xls(inhoud: bytes) -> Grid:
    import xlrd

    wb = xlrd.open_workbook(file_contents=inhoud)
    bladen: dict[str, Grid] = {}
    for sh in wb.sheets():
        rijen: list[dict[int, Cel]] = []
        for r in range(sh.nrows):
            rij: dict[int, Cel] = {}
            for c in range(sh.ncols):
                cel = sh.cell(r, c)
                if cel.ctype == xlrd.XL_CELL_EMPTY or cel.ctype == xlrd.XL_CELL_BLANK:
                    continue
                w = _naar_cel(cel.value, xls_datum=cel.ctype == xlrd.XL_CELL_DATE, datemode=wb.datemode)
                if w is not None:
                    rij[c] = w
            rijen.append(rij)
        bladen[sh.name] = Grid(naam=sh.name, rijen=rijen)
    eerste = next(iter(bladen.values()))
    eerste.bladen = bladen
    return eerste


def _lees_xlsx(inhoud: bytes) -> Grid:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(inhoud), data_only=True, read_only=True)
    bladen: dict[str, Grid] = {}
    for ws in wb.worksheets:
        rijen: list[dict[int, Cel]] = []
        for row in ws.iter_rows(values_only=True):
            rij: dict[int, Cel] = {}
            for c, w in enumerate(row):
                cel = _naar_cel(w)
                if cel is not None:
                    rij[c] = cel
            rijen.append(rij)
        bladen[ws.title] = Grid(naam=ws.title, rijen=rijen)
    eerste = next(iter(bladen.values()))
    eerste.bladen = bladen
    return eerste
